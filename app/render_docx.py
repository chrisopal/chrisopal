"""DOCX rendering utilities."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.utils.log import get_logger

logger = get_logger(__name__)


def _configure_styles(document: Document) -> None:
    style = document.styles["Normal"]
    font = style.font
    font.name = "SimSun"
    font.size = Pt(11)


def export_docx(sections: Iterable[dict], out_path: Path, title: str, brand: str) -> Path:
    """Render the proposal document to ``out_path`` and return the path."""

    document = Document()
    _configure_styles(document)

    title_para = document.add_heading(title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph(f"编制单位：{brand}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(12)

    info_para = document.add_paragraph(
        "编制日期：" + datetime.now().strftime("%Y年%m月%d日")
    )
    info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_page_break()

    for idx, section in enumerate(sections, start=1):
        heading = document.add_heading(f"第{idx}章 {section['title']}", level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for paragraph_text in section["content"].split("\n"):
            if not paragraph_text.strip():
                continue
            paragraph = document.add_paragraph(paragraph_text)
            paragraph.paragraph_format.space_after = Pt(6)
        document.add_paragraph("要点总结：" + section["content"].replace("\n", " "))
        document.add_paragraph(
            "本章节内容围绕需求、方案、保障等多角度展开，确保评审专家能够全面理解。"
        )
        document.add_page_break()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    logger.info("DOCX saved to %s", out_path)
    return out_path


__all__ = ["export_docx"]
